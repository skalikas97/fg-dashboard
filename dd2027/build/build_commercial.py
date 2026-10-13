"""AL-IDRISI DD2027 - commercial workbook (used when we perform the DD), built from spec.

Construction from SANTIAGO DOCKING / SO - SOW (build/spec/*), AA style wins (T8): Arial, Rev order
"Rev. # | Name | Date", AA header grey, Ongoing yellow, no tab colours, one workbook-wide revision
held on the master tab SOW initial. Money tabs: SANTIAGO DOCKING colours (T9).
Usage: python3 -I build/build_commercial.py out/AI_DD2027_Commercial_Rev00.xlsx
"""
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from openpyxl import Workbook
from openpyxl.formatting.rule import DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.formatting.rule import Rule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

import frame
from build_execution import top_block, arial, fill, border, style, CENTRE, GREY_HEADER
from common import CATEGORIES, COST_CARRIER, DATE_FMT, REV_ROWS, finish
import planner_controls

L = frame.L
F = "Arial"                              # T8: AA font everywhere
MASTER = "SOW initial"
BUFFER = 5                               # T4
COST = "Costfile - Invoice discussion"
SD = "SANTIAGO_DOCKING__"


def chapter(label):
    """'A. PILOTAGE, MOORING AND UNMOORING' -> 'Pilotage, mooring and unmooring'."""
    name = label.split(". ", 1)[1]
    return name[0] + name[1:].lower()


def latest_rev(k):
    rng = f"'{MASTER}'!${L(7 + k)}$3:${L(7 + k)}${2 + REV_ROWS}"
    return f'=IFERROR(LOOKUP(2,1/({rng}<>""),{rng}),"")'


def rev_compact(ws, c1, row):
    """Rev. # | Name | Date header, latest revision of the workbook, Last change (AA order)."""
    for k, txt in enumerate(("Rev. #", "Name", "Date")):
        style(ws.cell(row, c1 + k, txt), arial(10, bold=True), brd=border(
            l="medium" if k == 0 else "thin", r="medium" if k == 2 else "thin", t="medium", b="thin"), al=CENTRE)
        style(ws.cell(row + 1, c1 + k, latest_rev(k)), arial(10), brd=border(
            l="medium" if k == 0 else "thin", r="medium" if k == 2 else "thin", t="thin", b="thin"), al=CENTRE,
            nf=DATE_FMT if k == 2 else None)
    ws.merge_cells(start_row=row + 2, start_column=c1, end_row=row + 2, end_column=c1 + 1)
    style(ws.cell(row + 2, c1, "Last change"), arial(10), brd=border(l="medium", t="thin", b="medium"), al=CENTRE)
    ws.cell(row + 2, c1 + 1).border = border(r="thin", t="thin", b="medium")
    style(ws.cell(row + 2, c1 + 2, "=LastChange"), arial(10), brd=border(l="thin", r="medium", t="thin", b="medium"),
          al=CENTRE, nf=DATE_FMT)


def last_change_cell(ws, row, col):
    style(ws.cell(row, col, "Last change"), arial(10), al=Alignment(horizontal="right", vertical="center"))
    style(ws.cell(row, col + 1, "=LastChange"), arial(10), brd=border(b="thin"), al=CENTRE, nf=DATE_FMT)


def local_name(ws, name, ref):
    ws.defined_names[name] = DefinedName(name, attr_text=f"'{ws.title}'!{ref}")


# ---------------------------------------------------------------- SOW initial / SOW final
def build_sow(wb, title, subtitle, master):
    spec = frame.load("SO_-_SOW__SOW_final")
    ws = wb.create_sheet(title)
    ws.column_dimensions["A"].width = 2.0                   # AA spacer column
    frame.widths(spec, ws, 1, 11, offset=1)                 # SO columns A-K -> B-L
    legend = [(11, 11, 2 + i, t, rgb, 10, False) for i, (t, rgb) in enumerate(
        [("not yet started", None), ("Ongoing", "FFFF00"), ("Done", "92D050"), ("TBC", "FFC000"),
         ("Cancelled", "FF0000")])]
    last = top_block(ws, subtitle, 6, [(7, 7), (8, 8), (9, 9)], master, legend,
                     master_name=MASTER, rev_src=7)
    ws.merge_cells(start_row=2, start_column=12, end_row=last, end_column=12)   # SO K1:K5 grey block
    for r in range(2, last + 1):
        ws.cell(r, 12).fill = fill("A6A6A6")
    if master:
        wb.defined_names["LastChange"] = DefinedName("LastChange", attr_text=f"'{MASTER}'!$I${last}")
    hdr = last + 2
    ws.row_dimensions[hdr - 1].height = 15.0
    for c in range(1, 12):
        d = dict(frame.cell(spec, f"{L(c)}7"))
        d["fill"] = GREY_HEADER                              # AA header grey
        frame.apply(ws.cell(hdr, c + 1), d, F, value=True)
    ws.row_dimensions[hdr].height = 38.25
    ws.cell(hdr, 12, "Remarks")                              # reference header reads "Remakrs"
    r = hdr + 1
    for label in CATEGORIES:
        for c in range(1, 12):
            d = dict(frame.cell(spec, f"{L(c)}8"))
            d["font"] = [F, 10.0, True, False, None, None]
            frame.apply(ws.cell(r, c + 1), d, F)
        ws.cell(r, 2, label)
        r += 1
        for _ in range(BUFFER):
            for c in range(1, 12):
                d = dict(frame.cell(spec, f"{L(c)}9"))
                d["font"] = [F, 10.0, c == 5, False, None, None]     # Afloat/Dock letters bold
                frame.apply(ws.cell(r, c + 1), d, F, fill=False)
            r += 1
    end = r - 1
    ws.freeze_panes = ws.cell(hdr + 1, 1)
    ws.auto_filter.ref = f"B{hdr}:L{end}"
    frame.page(spec, ws, titles=f"{hdr}:{hdr}")
    ws.sheet_properties.tabColor = None
    ws.sheet_view.showGridLines = True                       # AA view
    return ws


# ---------------------------------------------------------------- money tabs
def fill_in(text, R, first, last):
    """Replace {R(n)} (reference row n moved to the new layout), {first}, {last} in a formula."""
    text = re.sub(r"\{R\((\d+)\)\}", lambda m: str(R(int(m.group(1)))), text)
    return text.replace("{first}", str(first)).replace("{last}", str(last))


def money_tab(wb, title, ref, cols, yard_col, rows, formulas, subtotal_cols, codes, bottom):
    """Costfile / Budget: top rows 1-7 as reference, per category BUFFER lines + Subtotal row,
    then the total block and the cost-code summary rebuilt with the new row numbers."""
    spec = frame.load(SD + ref)
    ws = wb.create_sheet(title)
    frame.widths(spec, ws, 1, cols)
    for r in range(1, 8):
        frame.copy_row(spec, r, ws, r, 1, cols, F, values=True)
    ws["A2"] = "PRICE COMPARISON - DRYDOCKING AL-IDRISI"
    ws["A3"], ws["B3"], ws["C3"] = "1 <CUR> = ", None, "Wisselkoers <date>"
    ws["B4"], ws["B5"] = None, None
    ws["B4"].number_format = "#,##0.0000"            # reference format carried the old local currency
    ws.cell(6, yard_col, "Yard 1")
    local_name(ws, "USD", "$B$4")
    local_name(ws, "EUR", "$B$5")

    first = 8
    r = first
    subtotals = []
    tmpl, sub_ref = rows
    for label in CATEGORIES:
        start = r
        for _ in range(BUFFER):
            frame.copy_row(spec, tmpl, ws, r, 1, cols, F)
            for col, f in formulas.items():
                ws[f"{col}{r}"] = f.format(r=r)
            r += 1
        frame.copy_row(spec, sub_ref, ws, r, 1, cols, F)
        ws.cell(r, 2, f"Subtotal {chapter(label)}")
        for col in subtotal_cols:
            ws[f"{col}{r}"] = f"=SUBTOTAL(9,{col}{start}:{col}{r - 1})"
        subtotals.append(r)
        r += 1
    last_body = r - 1
    off = r - bottom["first"]                      # reference bottom block shifted onto the new rows
    for ref_row in range(bottom["first"], bottom["last"] + 1):
        frame.copy_row(spec, ref_row, ws, ref_row + off, 1, cols + 2, F)
    def R(ref_row):
        return ref_row + off
    for (ref_row, col), val in bottom["cells"].items():
        ws[f"{col}{R(ref_row)}"] = fill_in(val, R, first, last_body)
    # cost-code summary
    code_col, desc_col, amt_col, disc_col = bottom["summary_cols"]
    key_col, val_col = bottom["summary_keys"]
    rate = fill_in(bottom["summary_rate"], R, first, last_body)
    sr = R(bottom["summary_first"])
    for i, (code, desc) in enumerate(codes):
        row = sr + i if code != "Project" else R(bottom["project_row"])
        ws[f"{code_col}{row}"] = f"0080/{COST_CARRIER}/{code}" if code != "Project" else "Project"
        ws[f"{desc_col}{row}"] = desc
        ws[f"{amt_col}{row}"] = (f"=SUMPRODUCT((${key_col}${first}:${key_col}${last_body}={code_col}{row})"
                                 f"*(${val_col}${first}:${val_col}${last_body}))")
        ws[f"{disc_col}{row}"] = f"={amt_col}{row}*(1-{rate})"
    tot = R(bottom["total_row"])
    ws[f"{amt_col}{tot}"] = f"=SUM({amt_col}{sr}:{amt_col}{tot - 2})"
    ws[f"{disc_col}{tot}"] = f"=SUM({disc_col}{sr}:{disc_col}{tot - 2})"
    ws.freeze_panes = "A7"
    frame.page(spec, ws, area=f"A2:{L(bottom.get('area_col', cols))}{tot}")
    ws.sheet_properties.tabColor = None
    return ws, R, last_body, subtotals


COST_CODES = [("400", "O&R Main Drives"), ("401", "O&R  Aux Drives"), ("402", "O&R Electrical"),
              ("403", "O&R Hydraulic"), ("404", "O&R Hoisting Equipment"), ("405", "O&R Engine Room"),
              ("406", "O&R Dredging Equipment"), ("407", "O&R Hull"), ("408", "O&R Hopper"),
              ("411", "O&R General")]


def build_costfile(wb):
    codes = COST_CODES + [("447", "Shipyard repairs"), ("449", "M&R Safety, Classification & Inspection"),
                          ("453", "Fuel"), ("463", "Fresh Water"), ("464", "Waste Disposal"),
                          ("484", "Shore power supply"), ("Project", "Project Logistics")]
    bottom = {
        "first": 516, "last": 547,
        "cells": {
            (518, "F"): "=SUBTOTAL(9,F{first}:F{last})", (518, "H"): "=SUBTOTAL(9,H{first}:H{last})",
            (518, "I"): "=SUBTOTAL(9,I{first}:I{last})", (518, "L"): "=SUBTOTAL(9,L{first}:L{last})",
            (518, "M"): "=SUBTOTAL(9,M{first}:M{last})", (518, "N"): "=SUBTOTAL(9,N{first}:N{last})",
            (518, "O"): "=SUBTOTAL(9,O{first}:O{last})", (518, "P"): "=SUBTOTAL(9,P{first}:P{last})",
            (518, "Q"): "=P{R(518)}*EUR",
            (519, "B"): "TOTAL WITHOUT DISCOUNT", (519, "E"): "=F{R(518)}*EUR",
            (520, "B"): "DISCOUNT", (520, "F"): "=F{R(518)}*E{R(520)}", (520, "I"): "=I{R(518)}*H{R(520)}",
            (520, "P"): "=P{R(518)}*O{R(520)}", (520, "Q"): "=P{R(520)}*EUR",
            (521, "E"): "=F{R(520)}*EUR",
            (522, "B"): "GRAND TOTAL", (522, "F"): "=F{R(518)}-F{R(520)}", (522, "I"): "=I{R(518)}-I{R(520)}",
            (522, "P"): "=P{R(518)}-P{R(520)}",
            (523, "E"): "=F{R(522)}*EUR", (523, "H"): "=I{R(522)}*EUR", (523, "O"): "Already paid",
            (523, "Q"): "=P{R(523)}*EUR",
            (524, "B"): "Completion", (524, "O"): "Remaining", (524, "P"): "=P{R(522)}-P{R(523)}",
            (524, "Q"): "=P{R(524)}*EUR",
            (527, "L"): "CODE", (527, "N"): "Description", (527, "O"): "Amount", (527, "P"): "Amount With discount",
            (547, "M"): "TOTAL",
        },
        "summary_cols": ("L", "N", "O", "P"), "summary_keys": ("R", "P"), "summary_rate": "$O${R(520)}",
        "summary_first": 528, "project_row": 545, "total_row": 547,
    }
    formulas = {"F": "=D{r}*E{r}", "I": "=H{r}*G{r}", "P": "=I{r}-O{r}"}
    ws, R, last_body, subs = money_tab(wb, COST, "Costfile_-_Invoice_discussion", 18, 5, (11, 21),
                                        formulas, "FHILMNOP", codes, bottom)
    grand = R(522)
    for r in range(8, last_body + 1):
        ws[f"J{r}"] = f"=IFERROR(I{r}/$I${grand},0)"
        if r not in subs:
            ws[f"I{r}"].fill = PatternFill()          # status colour on Draft Invoice is set by the TS
    for r in subs:
        ws[f"I{r}"].fill = fill("BFBFBF")
    ws["I2"], ws["I3"], ws["I4"], ws["I5"] = "not completed yet", "agreed", "not agreed", "discussion"
    ws["P6"] = "Final"
    for c in "PQR":
        for r in (2, 3, 4):
            ws[f"{c}{r}"] = None
    rev_compact(ws, 16, 2)
    ws.auto_filter.ref = f"R8:R{last_body}"
    return R


def build_budget(wb):
    codes = COST_CODES + [("453", "Fuel"), ("463", "Fresh Water"), ("464", "Waste Disposal"),
                          ("Project", "Project Logistics")]
    bottom = {
        "first": 522, "last": 550,
        "cells": {
            (525, "F"): "=SUBTOTAL(9,F{first}:F{last})", (525, "G"): "=SUBTOTAL(9,G{first}:G{last})",
            (525, "H"): "=SUBTOTAL(9,H{first}:H{last})", (525, "I"): "=SUBTOTAL(9,I{first}:I{last})",
            (525, "J"): "=SUBTOTAL(9,J{first}:J{last})", (525, "K"): "=SUBTOTAL(9,K{first}:K{last})",
            (525, "L"): "=SUBTOTAL(9,L{first}:L{last})", (525, "M"): "=SUBTOTAL(9,M{first}:M{last})",
            (526, "B"): "TOTAL WITHOUT DISCOUNT", (526, "E"): "=F{R(525)}*EUR",
            (527, "B"): "DISCOUNT", (527, "F"): "=F{R(525)}*E{R(527)}", (527, "H"): "=H{R(525)}*G{R(527)}",
            (527, "M"): "=M{R(525)}*L{R(527)}",
            (528, "E"): "=F{R(527)}*EUR",
            (529, "B"): "GRAND TOTAL", (529, "F"): "=F{R(525)}-F{R(527)}", (529, "H"): "=H{R(525)}-H{R(527)}",
            (529, "M"): "=M{R(525)}-M{R(527)}",
            (530, "E"): "=F{R(529)}*EUR", (530, "L"): "Already paid",
            (531, "B"): "Completion", (531, "L"): "Remaining", (531, "M"): "=M{R(529)}-M{R(530)}",
            (534, "I"): "CODE", (534, "J"): "Description", (534, "L"): "Amount", (534, "M"): "Amount With discount",
            (550, "K"): "TOTAL",
        },
        "summary_cols": ("I", "J", "L", "M"), "summary_keys": ("O", "M"), "summary_rate": "$L${R(527)}",
        "summary_first": 535, "project_row": 548, "total_row": 550, "area_col": 15,
    }
    ws, R, last_body, subs = money_tab(wb, "Budget file", "Budget_file", 16, 5, (10, 511),
                                        {"F": "=D{r}*E{r}"}, "FGHIJKLM", codes, bottom)
    rev_compact(ws, 13, 2)
    return ws


# ---------------------------------------------------------------- AXI
def build_axi(wb, cost_R):
    spec = frame.load(SD + "AXI_cost_breakdown")
    ws = wb.create_sheet("AXI cost breakdown")
    frame.widths(spec, ws, 1, 15)
    for r in range(1, 82):
        frame.copy_row(spec, r, ws, r, 1, 15, F, values=True)
    ws.merge_cells("A1:L4")
    ws.merge_cells("A42:D42")
    ws["A1"] = "AL-IDRISI DRYDOCK 2027 - AXI COST BREAKDOWN"
    for c in "MNO":
        for r in (1, 2, 3):
            ws[f"{c}{r}"] = None
    rev_compact(ws, 13, 1)
    # summary rows of the Costfile per cost code (same workbook - no external link)
    code_row = {c: cost_R(528) + i for i, c in enumerate(
        ["400", "401", "402", "403", "404", "405", "406", "407", "408", "411", "447", "449", "453", "463", "464", "484"])}
    for r in range(6, 38):
        code = ws[f"C{r}"].value
        g = ws[f"G{r}"].value
        if isinstance(g, str) and "[1]" in g:
            ws[f"G{r}"] = f"='{COST}'!O{code_row[str(code)]}"
        if ws[f"B{r}"].value == 8579:
            ws[f"B{r}"] = int(COST_CARRIER)
        for col in "HIJKL":
            v = ws[f"{col}{r}"].value
            if isinstance(v, str) and v.startswith("=") and r < 38:
                ws[f"{col}{r}"] = f"=IFERROR({v[1:]},0)"
    ws["H38"] = f"='{COST}'!P{cost_R(520)}"
    ws["L38"] = f"='{COST}'!H{cost_R(523)}"
    ws.auto_filter.ref = "A5:L5"
    frame.page(spec, ws)
    ws.sheet_properties.tabColor = None


# ---------------------------------------------------------------- Scaffolding
def build_scaffolding(wb):
    spec = frame.load(SD + "Scaffolding")
    ws = wb.create_sheet("Scaffolding")
    frame.widths(spec, ws, 1, 16)
    for r in (2, 3):
        frame.copy_row(spec, r, ws, r, 1, 16, F, values=True)
    for r in range(4, 34):
        frame.copy_row(spec, 4, ws, r, 1, 16, F)
        ws[f"H{r}"].fill = PatternFill()
        ws[f"E{r}"] = f'=IF(C{r}="","",C{r}*D{r})'
        ws[f"N{r}"] = f'=IF(L{r}="","",L{r}*M{r})'
    frame.copy_row(spec, 34, ws, 34, 1, 16, F)
    ws["N34"], ws["O34"] = "=SUM(N4:N33)", "=SUM(O4:O33)"
    for r, (txt, letter) in zip((37, 38, 39), (("AGREED =", "A"), ("DISCUSSION =", "D"), ("NOT AGREED = ", "N"))):
        frame.copy_row(spec, r, ws, r, 1, 16, F)
        ws[f"B{r}"], ws[f"D{r}"] = txt, letter
        ws[f"C{r}"] = f'=SUMIF($H$4:$H$33,"{letter}",$N$4:$N$33)'
    ws.merge_cells("J37:K37")
    ws.merge_cells("J7:K7")
    last_change_cell(ws, 1, 13)
    frame.page(spec, ws)
    ws.sheet_properties.tabColor = None


# ---------------------------------------------------------------- 3rd parties
def build_third_parties(wb, vendors=6):
    spec = frame.load(SD + "3rd_parties")
    ws = wb.create_sheet("3rd parties")
    frame.widths(spec, ws, 1, 17)
    frame.copy_row(spec, 1, ws, 1, 1, 12, F, values=True)
    ws.merge_cells("A1:L1")
    rev_compact(ws, 13, 1)
    lines = ["Air Travel", "Hotel", "Transport", "Working time"]
    r = 2
    for n in range(1, vendors + 1):
        frame.copy_row(spec, 2, ws, r, 1, 12, F)
        ws[f"A{r}"], ws[f"B{r}"] = f"{n}) ", "<VENDOR> (<SCOPE>)"
        frame.copy_row(spec, 3, ws, r + 1, 1, 12, F, values=True)
        for i in range(5):                       # 4 fixed lines + 1 empty line (ref rows 4-8)
            rr = r + 2 + i
            frame.copy_row(spec, 4 + i, ws, rr, 1, 12, F)
            if i < 4:
                ws[f"A{rr}"], ws[f"B{rr}"] = i + 1, lines[i]
        ws.merge_cells(start_row=r + 2, start_column=11, end_row=r + 5, end_column=11)
        ws.merge_cells(start_row=r + 2, start_column=12, end_row=r + 5, end_column=12)
        frame.copy_row(spec, 9, ws, r + 7, 1, 12, F, values=True)
        ws[f"F{r + 7}"] = f"=SUM(F{r + 2}:F{r + 6})"
        frame.copy_row(spec, 10, ws, r + 8, 1, 12, F)
        r += 9
    frame.page(spec, ws)
    ws.sheet_properties.tabColor = None


# ---------------------------------------------------------------- 3rd party Planner
def build_planner(wb, rows=37, weeks=8):
    spec = frame.load(SD + "3rd_party_Planner")
    ws = wb.create_sheet("3rd party Planner")
    g1, ndays = 11, weeks * 7                       # K .. BN
    gl = g1 + ndays - 1
    frame.widths(spec, ws, 1, 10)
    for c in range(g1, gl + 1):
        ws.column_dimensions[L(c)].width = 2.42578125
    for r in range(1, 8):
        frame.copy_row(spec, r, ws, r, 1, 10, F, values=True)
        for c in range(g1, gl + 1):
            ref = L(g1 + (c - g1) % 7 if (c - g1) >= 7 else c)
            frame.apply(ws.cell(r, c), frame.cell(spec, f"{ref}{r}"), F)
    ws["A1"] = "AL-IDRISI - Planner"
    ws["C4"] = None
    ws["C5"] = "=TODAY()"
    for c in ("C4", "C5"):
        ws[c].number_format = DATE_FMT
    ws.merge_cells("C4:D4")
    ws.merge_cells("C5:D5")
    ws.merge_cells(start_row=1, start_column=g1, end_row=1, end_column=g1 + 20)
    last_change_cell(ws, 2, 2)
    ws["C2"].border = border(b="thin")
    for w in range(weeks):
        c = g1 + 7 * w
        ws.cell(4, c, f'=IF($C$4="","","Week "&({L(c)}6-($C$4-WEEKDAY($C$4,1)+2))/7+1)')
        ws.cell(5, c, f"={L(c)}6")
        ws.merge_cells(start_row=4, start_column=c, end_row=4, end_column=c + 6)
        ws.merge_cells(start_row=5, start_column=c, end_row=5, end_column=c + 6)
    for c in range(g1, gl + 1):
        ws.cell(6, c, '=IF($C$4="","",$C$4-WEEKDAY($C$4,1)+2+7*($G$4-1))' if c == g1
                else f'=IF({L(c - 1)}6="","",{L(c - 1)}6+1)')
        ws.cell(7, c, f'=IF({L(c)}6="","",CHOOSE(WEEKDAY({L(c)}6,1),"S","M","T","W","T","F","S"))')
    dv = DataValidation(allow_blank=True, showInputMessage=True, showErrorMessage=False,
                        promptTitle="Display Week",
                        prompt="Enter the week number to display first in the Gantt Chart. The weeks are numbered "
                               "starting from the week containing the Project Start Date.")
    dv.add("G4")
    ws.add_data_validation(dv)

    frame.copy_row(spec, 8, ws, 8, 1, 10, F, values=True)
    for c in range(g1, gl + 1):
        frame.apply(ws.cell(8, c), frame.cell(spec, "K8"), F)
    ws["E8"] = None
    ws.defined_names["prevWBS"] = DefinedName("prevWBS", attr_text=f"'{ws.title}'!$A1048576")
    for r in range(9, rows + 1):
        frame.copy_row(spec, 9, ws, r, 1, 10, F)
        for c in range(g1, gl + 1):
            frame.apply(ws.cell(r, c), frame.cell(spec, "K9"), F)
        ws[f"F{r}"] = (f'=IF(D{r}="","",IF(D{r}="N/A","N/A",IF(($C$5-D{r}+1)/G{r}>1,1,IF(($C$5-D{r}+1)/G{r}<0,0,'
                       f'IF(OR(D{r}=0,$C$5=0),"-",($C$5-D{r}+1)/G{r})))))')
        ws[f"G{r}"] = f'=IF(D{r}="","",IF(D{r}="N/A","N/A",IF(OR(E{r}=0,D{r}=0)," - ",NETWORKDAYS(D{r},E{r}))))'
        ws[f"I{r}"] = f'=IF(D{r}="","",IF(AND(H{r}="-",$C$5>E{r}),$C$5-E{r},"-"))'
        for c in "DEH":
            ws[f"{c}{r}"].number_format = DATE_FMT
    G = f"{L(g1)}"
    grid = f"{G}9:{L(gl)}{rows}"
    ws.conditional_formatting.add(f"{G}6:{L(gl)}{rows}", Rule(type="expression", formula=[f"{G}$6=TODAY()"],
        dxf=DifferentialStyle(border=Border(left=Side("thin", color="FFC00000"), right=Side("thin", color="FFC00000")))))
    ws.conditional_formatting.add(f"{G}6:{L(gl)}7", FormulaRule(formula=[f"{G}$6=TODAY()"], fill=fill("C0504D"),
                                                               font=Font(color="FFFFFFFF")))
    for f, colour in ((f'AND(NOT(ISBLANK($E9)),$E9<{G}$6,$C$5>={G}$6,$H9="-")', "FF0000"),
                      (f'AND($D9<>"",{G}$6<>"",OR(AND($H9<>"-",$D9<={G}$6,$H9>={G}$6),AND($D9<={G}$6,$E9>={G}$6,$F9=1,NOT($H9="-"))))', "92D050"),
                      (f'AND($D9<>"",$H9="-",$D9<={G}$6,ROUNDDOWN(($E9-$D9+1)*$F9,0)+$D9-1>={G}$6)', "FFC000"),
                      (f'AND($H9="-",NOT(ISBLANK($D9)),$D9<={G}$6,$E9>={G}$6)', "0070C0")):
        ws.conditional_formatting.add(grid, FormulaRule(formula=[f], fill=fill(colour)))
    ws.conditional_formatting.add(f"F8:F{rows}", DataBarRule(start_type="num", start_value=0, end_type="num",
                                                              end_value=1, color="FFFFC000"))
    ws.freeze_panes = "A8"
    frame.page(spec, ws, area=f"A1:{L(gl)}{rows}", titles="4:7")
    ws.sheet_properties.tabColor = None
    return ws


ORDER = ["Costfile - Invoice discussion", "AXI cost breakdown", "Scaffolding", "3rd party Planner",
         "3rd parties", "SOW final", "SOW initial", "Budget file"]


def main(out):
    wb = Workbook()
    wb.remove(wb.active)
    build_sow(wb, "SOW initial", "Estimated scope of works", True)
    build_sow(wb, "SOW final", "Actual scope of works", False)
    cost_R = build_costfile(wb)
    build_budget(wb)
    build_axi(wb, cost_R)
    build_scaffolding(wb)
    build_third_parties(wb)
    planner = build_planner(wb)
    wb._sheets = [wb[n] for n in ORDER]
    wb.active = ORDER.index(MASTER)
    idx = ORDER.index(planner.title) + 1
    finish(wb, out, patch=lambda files: planner_controls.inject(files, idx))


if __name__ == "__main__":
    main(sys.argv[1])
