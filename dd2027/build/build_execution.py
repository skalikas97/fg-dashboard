"""AL-IDRISI DD2027 - execution workbook (AA style), built from spec.

Spec: pack/03 A1, A4, A7 and pack/03b "AA :: Daily SOW", plus TS decisions in STATE.md.
One function per tab; tabs are added stage by stage after TS sign-off.
Usage: python3 -I build/build_execution.py out/AI_DD2027_Execution_Rev00.xlsx
"""
import os
import sys
import zipfile
from datetime import datetime, date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.properties import PageSetupProperties

AUTHOR = "Psyroukis Konstantinos"
VESSEL_TITLE = "AL-IDRISI - Drydock 2027 <Place>"
REV0 = ("Rev00", "PSKO", date(2026, 10, 13))
REV_ROWS = 10        # T4: 10 pre-formatted Rev rows
BUFFER_ROWS = 5      # T4: 5 empty rows per category (15 on Spares)

# Yard tariff letters (T3), verbatim from SO - SOW.xlsx, SOW initial / SOW final
CATEGORIES = [
    "A. PILOTAGE, MOORING AND UNMOORING", "B. DRYDOCKING", "C. SERVICES",
    "D. GENERAL", "E. SURFACE PREPARATIONS", "F. PAINT", "G. STEEL", "H. TANKS",
    "I. SEAWATER LINES", "J. HOPPER", "M. PROPELLERS AND TAILSHAFTS",
    "N. ENGINES", "Q. AUXILIARIES", "T. ANCHORS AND CHAINS", "X. MISCELLANEOUS",
]

# AA palette as plain RGB (pack/03 A4, AA theme resolved)
GREY_HEADER = "A6A6A6"   # th0-0.35
GREY_CAT = "BFBFBF"      # th0-0.25
TITLE_RED = "A80000"
TITLE_FONT = "F2F2F2"    # th0-0.05
LEGEND = [("Planned", None), ("Ongoing", "FFFF00"), ("Completed", "92D050"),
          ("TBC", "FFC000"), ("Cancelled/postponed", "FF0000"),
          ("All necessary parts on board", "00B0F0")]

DATE_FMT = "dd/mm/yyyy"


def side(style):
    return Side(style=style) if style else Side()


def border(l=None, r=None, t=None, b=None):
    return Border(left=side(l), right=side(r), top=side(t), bottom=side(b))


def fill(rgb):
    return PatternFill("solid", fgColor="FF" + rgb) if rgb else PatternFill()


def arial(size, bold=False, color=None):
    return Font(name="Arial", size=size, bold=bold, color=("FF" + color) if color else None)


def style(cell, font=None, fil=None, brd=None, al=None, nf=None):
    if font: cell.font = font
    if fil: cell.fill = fil
    if brd: cell.border = brd
    if al: cell.alignment = al
    if nf: cell.number_format = nf


def top_block(ws, subtitle, master):
    """AA top block: spacer row 1, title B2:E<last>, Rev block F:H, Last change row, legend J."""
    last = 2 + REV_ROWS + 1                      # Rev header + rev rows + Last change row
    ws.row_dimensions[1].height = 10.5
    ws.row_dimensions[2].height = 15.0
    for r in range(3, last):
        ws.row_dimensions[r].height = 14.25
    ws.row_dimensions[last].height = 15.0

    # Title (merged), borders on the edge cells of the merge
    ws.merge_cells(start_row=2, start_column=2, end_row=last, end_column=5)
    t = ws.cell(2, 2, f"{VESSEL_TITLE}\n{subtitle}")
    style(t, arial(28, color=TITLE_FONT), fill(TITLE_RED),
          al=Alignment(horizontal="center", vertical="center", wrap_text=True))
    for r in range(2, last + 1):
        for c in range(2, 6):
            ws.cell(r, c).border = border(l="medium" if c == 2 else None,
                                          t="medium" if r == 2 else None,
                                          b="medium" if r == last else None)

    # Rev block: "Rev. # | Name | Date" (AA order, T8)
    centre = Alignment(horizontal="center", vertical="center")
    for c, txt in zip((6, 7, 8), ("Rev. #", "Name", "Date")):
        style(ws.cell(2, c, txt), arial(10), al=centre,
              brd=border(l="medium" if c == 6 else "thin" if c == 7 else None,
                         r="medium" if c == 8 else "thin", t="medium"))
    for i in range(REV_ROWS):
        r = 3 + i
        for c in (6, 7, 8):
            style(ws.cell(r, c), arial(10), al=centre,
                  nf=DATE_FMT if c == 8 else None,
                  brd=border(l="medium" if c == 6 else "thin",
                             r="medium" if c == 8 else "thin", t="thin", b="thin"))
        if master and i == 0:
            ws.cell(r, 6, REV0[0]); ws.cell(r, 7, REV0[1]); ws.cell(r, 8, REV0[2])
        elif not master:
            # secondary tabs pull the Rev block from the master tab (AA practice)
            for c in (6, 7, 8):
                ref = f"'Daily SOW'!{ws.cell(r, c).coordinate}"
                ws.cell(r, c, f'=IF({ref}="","",{ref})')

    # Last change row (T5/S1): named cell LastChange on the master, formula elsewhere
    ws.merge_cells(start_row=last, start_column=6, end_row=last, end_column=7)
    style(ws.cell(last, 6, "Last change"), arial(10), al=centre,
          brd=border(l="medium", t="medium", b="medium"))
    ws.cell(last, 7).border = border(r="thin", t="medium", b="medium")
    lc = ws.cell(last, 8, REV0[2] if master else "=LastChange")
    style(lc, arial(10), al=centre, nf=DATE_FMT,
          brd=border(l="thin", r="medium", t="medium", b="medium"))

    # Legend (no borders)
    for i, (txt, rgb) in enumerate(LEGEND):
        style(ws.cell(2 + i, 10, txt), arial(11), fill(rgb),
              al=Alignment(vertical="center"))
    return last


def build_daily_sow(wb):
    ws = wb.active
    ws.title = "Daily SOW"
    widths = {"A": 2.0, "B": 14.28515625, "C": 70.85546875, "D": 6.0, "E": 17.5703125,
              "F": 12.140625, "G": 12.42578125, "H": 16.42578125, "I": 62.42578125, "J": 28.0}
    for k, w in widths.items():
        ws.column_dimensions[k].width = w

    last = top_block(ws, "Scope of works", master=True)
    wb.defined_names["LastChange"] = DefinedName("LastChange", attr_text=f"'Daily SOW'!$H${last}")

    hdr = last + 2                                 # one blank row, then the header
    ws.row_dimensions[hdr - 1].height = 15.0
    ws.row_dimensions[hdr].height = 18.75
    heads = {2: ("Specs", "left", True, ("medium", "medium")),
             3: ("Description", "center", False, (None, "medium")),
             4: (None, "center", False, (None, "medium")),
             5: ("WL item", "center", True, ("medium", "medium")),
             6: ("3rd party", "center", True, ("medium", "medium")),
             7: ("Crew", "center", True, ("medium", "medium")),
             8: ("JDN WS", "center", True, ("medium", "medium")),
             9: ("Remarks", "left", True, ("medium", "medium"))}
    for c, (txt, h, wrap, (l, r)) in heads.items():
        style(ws.cell(hdr, c, txt), arial(14), fill(GREY_HEADER),
              border(l=l, r=r, t="medium", b="medium"),
              Alignment(horizontal=h, vertical="center", wrap_text=wrap or None))

    # Category rows + empty pre-formatted job rows
    cat_al = {c: Alignment(horizontal="left") for c in (2, 3, 4)}
    cat_al.update({c: Alignment(horizontal="center", vertical="center") for c in (5, 6, 7, 8)})
    cat_al[9] = Alignment(horizontal="left", wrap_text=True)
    data = {2: (border(l="medium", r="thin", t="dashed", b="dashed"), Alignment(horizontal="left", vertical="center")),
            3: (border(r="medium", t="dashed", b="dashed"), Alignment(vertical="center")),
            4: (border(r="medium", t="dashed", b="dashed"), Alignment(vertical="center")),
            5: (border(l="medium", r="medium", t="dashed", b="dashed"), Alignment(horizontal="center", vertical="center")),
            6: (border(l="medium", r="medium", t="dashed", b="dashed"), Alignment(horizontal="center", vertical="center")),
            7: (border(l="medium", r="medium", t="dashed", b="dashed"), Alignment(horizontal="center", vertical="center")),
            8: (border(l="medium", r="medium", t="dashed", b="dashed"), Alignment(horizontal="center", vertical="center")),
            9: (border(l="medium", r="medium", t="dashed", b="dashed"), Alignment(horizontal="left", vertical="center", wrap_text=True))}

    def cat_row(r, label):
        ws.row_dimensions[r].height = 15.0
        for c in range(2, 10):
            style(ws.cell(r, c), arial(10, bold=True), fill(GREY_CAT),
                  border(l="medium" if c == 2 else None, t="medium", b="medium"), cat_al[c])
        if label:
            ws.cell(r, 2, label)

    r = hdr + 1
    for label in CATEGORIES:
        cat_row(r, label)
        r += 1
        for _ in range(BUFFER_ROWS):
            for c, (brd, al) in data.items():
                style(ws.cell(r, c), arial(10), brd=brd, al=al)
            r += 1
    cat_row(r, None)                               # closing grey row, as in AA
    end = r

    ws.freeze_panes = ws.cell(hdr + 1, 1)
    ws.auto_filter.ref = f"A{hdr}:J{end}"
    ws.sheet_view.zoomScale = 85
    ws.print_area = f"A1:I{end}"
    ps = ws.page_setup
    ps.orientation, ps.paperSize, ps.fitToWidth, ps.scale = "portrait", 8, 0, 80
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    m = ws.page_margins
    m.left = m.right = 0.23622047244094488
    m.top = m.bottom = 0.7480314960629921
    m.header = m.footer = 0.31496062992125984
    return ws


def main(out):
    wb = Workbook()
    build_daily_sow(wb)
    now = datetime.now().replace(microsecond=0)
    p = wb.properties
    p.creator = p.lastModifiedBy = AUTHOR
    p.created = p.modified = now
    p.title = p.subject = p.description = p.keywords = p.category = None
    wb.save(out)
    strip_tool_names(out)


def strip_tool_names(path):
    """Blank docProps/app.xml so no library name is left in the file (rule: no tool names)."""
    app = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
           '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"/>')
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "docProps/app.xml":
                data = app.encode()
            dst.writestr(item, data)
    os.replace(tmp, path)


if __name__ == "__main__":
    main(sys.argv[1])
