"""Read-only extraction of the construction of the reference tabs into JSON specs.

Theme colours are resolved to plain RGB with each workbook's own theme (pack/03 A4).
Usage: python3 -I build/extract_spec.py <xlsx> <outdir> [maxrow] [maxcol]
Never saves the reference workbook.
"""
import colorsys
import json
import os
import re
import sys
import zipfile

import openpyxl
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.worksheet.formula import ArrayFormula


def theme_colours(path):
    xml = zipfile.ZipFile(path).read("xl/theme/theme1.xml").decode("utf8")
    scheme = re.search(r"<a:clrScheme.*?</a:clrScheme>", xml, re.S).group(0)
    out = {}
    for tag in ("dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3",
                "accent4", "accent5", "accent6"):
        block = re.search(rf"<a:{tag}>(.*?)</a:{tag}>", scheme, re.S).group(1)
        m = re.search(r'lastClr="([0-9A-Fa-f]{6})"', block) or re.search(r'val="([0-9A-Fa-f]{6})"', block)
        out[tag] = m.group(1).upper()
    order = ["lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6"]
    return [out[t] for t in order]


def apply_tint(rgb, tint):
    r, g, b = (int(rgb[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = l * (1 + tint) if tint < 0 else l * (1 - tint) + tint
    r, g, b = colorsys.hls_to_rgb(h, l, s)
    return "".join(f"{round(x * 255):02X}" for x in (r, g, b))


class Resolver:
    def __init__(self, path):
        self.theme = theme_colours(path)

    def __call__(self, c):
        if c is None:
            return None
        if c.type == "rgb":
            v = c.rgb if isinstance(c.rgb, str) else None
            return v[-6:].upper() if v else None
        if c.type == "theme":
            return apply_tint(self.theme[c.theme], c.tint or 0) if c.theme < 10 else None
        if c.type == "indexed":
            if c.indexed in (64, 65):          # system foreground / background
                return None
            return COLOR_INDEX[c.indexed][-6:] if c.indexed < len(COLOR_INDEX) else None
        return None


def side(s, res):
    return [s.style, res(s.color)] if s is not None and s.style else None


def cell_spec(c, res):
    f, fi, b, a = c.font, c.fill, c.border, c.alignment
    d = {}
    if isinstance(c.value, ArrayFormula):
        d["v"], d["array"] = c.value.text, c.value.ref
    elif c.value is not None:
        d["v"] = c.value if isinstance(c.value, (int, float, str)) else str(c.value)
        if hasattr(c.value, "isoformat"):
            d["date"] = True
    d["font"] = [f.name, f.sz, bool(f.b), bool(f.i), f.u, res(f.color)]
    if fi is not None and fi.fill_type == "solid":
        d["fill"] = res(fi.fgColor)
    br = {k: side(getattr(b, k), res) for k in ("left", "right", "top", "bottom")}
    br = {k: v for k, v in br.items() if v}
    if br:
        d["brd"] = br
    al = [a.horizontal, a.vertical, bool(a.wrap_text), a.text_rotation or 0, a.indent or 0]
    if any(al[:2]) or al[2] or al[3] or al[4]:
        d["al"] = al
    if c.number_format != "General":
        d["nf"] = c.number_format
    return d


def cf_rules(ws, res):
    out = []
    for rng, rules in ws.conditional_formatting._cf_rules.items():
        for r in rules:
            dx = r.dxf
            out.append({
                "range": str(rng.sqref), "type": r.type, "operator": r.operator, "text": r.text,
                "formula": list(r.formula or []), "stop": r.stopIfTrue, "priority": r.priority,
                "fill": res(dx.fill.bgColor) or res(dx.fill.fgColor) if dx and dx.fill else None,
                "font": res(dx.font.color) if dx and dx.font and dx.font.color else None,
                "bold": bool(dx.font.b) if dx and dx.font else False,
            })
    return out


def main(path, outdir, maxrow=800, maxcol=45):
    os.makedirs(outdir, exist_ok=True)
    res = Resolver(path)
    wb = openpyxl.load_workbook(path)
    tag = os.path.splitext(os.path.basename(path))[0]
    for ws in wb.worksheets:
        lastr = min(ws.max_row, maxrow)
        lastc = min(ws.max_column, maxcol)
        cells = {}
        for row in ws.iter_rows(min_row=1, max_row=lastr, max_col=lastc):
            for c in row:
                d = cell_spec(c, res)
                if set(d) - {"font"} or d["font"][0] not in (None, "Calibri", "Arial", "Aptos Narrow"):
                    cells[c.coordinate] = d
        pane = ws.sheet_view.pane
        ps, pm, hf = ws.page_setup, ws.page_margins, ws.oddFooter
        spec = {
            "sheet": ws.title, "tab": res(ws.sheet_properties.tabColor) if ws.sheet_properties.tabColor else None,
            "rows": {r: [d.height, d.hidden, d.outlineLevel] for r, d in ws.row_dimensions.items() if r <= lastr and (d.height or d.hidden or d.outlineLevel)},
            "cols": {k: [d.width, d.hidden, d.outlineLevel, d.min, d.max] for k, d in ws.column_dimensions.items()},
            "merges": [str(m) for m in ws.merged_cells.ranges if m.min_row <= lastr],
            "cells": cells, "cf": cf_rules(ws, res),
            "dv": [{"sqref": str(v.sqref), "type": v.type, "f1": v.formula1, "f2": v.formula2, "op": v.operator,
                    "prompt": v.prompt, "ptitle": v.promptTitle, "error": v.error, "etitle": v.errorTitle,
                    "showin": v.showInputMessage, "showerr": v.showErrorMessage, "blank": v.allow_blank}
                   for v in ws.data_validations.dataValidation],
            "freeze": [pane.ySplit if pane else 0, pane.xSplit if pane else 0],
            "filter": ws.auto_filter.ref, "zoom": ws.sheet_view.zoomScale, "grid": ws.sheet_view.showGridLines,
            "print": [ps.orientation, ps.paperSize, ps.fitToWidth, ps.fitToHeight, ps.scale,
                      bool(ws.sheet_properties.pageSetUpPr and ws.sheet_properties.pageSetUpPr.fitToPage)],
            "margins": [pm.left, pm.right, pm.top, pm.bottom, pm.header, pm.footer],
            "footer": [hf.left.text, hf.center.text, hf.right.text] if hf else None,
            "titles": ws.print_title_rows, "area": ws.print_area,
            "defrow": ws.sheet_format.defaultRowHeight, "summary_below": ws.sheet_properties.outlinePr.summaryBelow if ws.sheet_properties.outlinePr else None,
            "protect": bool(ws.protection.sheet),
        }
        name = f"{tag}__{ws.title}".replace(" ", "_").replace("/", "-")
        with open(os.path.join(outdir, name + ".json"), "w") as fh:
            json.dump(spec, fh, default=str)
        print(name, len(cells), "cells", lastr, "x", lastc)


if __name__ == "__main__":
    a = sys.argv
    main(a[1], a[2], int(a[3]) if len(a) > 3 else 800, int(a[4]) if len(a) > 4 else 45)
