"""Shared helpers: apply the reference construction (JSON specs from extract_spec.py) to new sheets."""
import json
import os

from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import column_index_from_string, coordinate_to_tuple, get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

SPEC_DIR = os.path.join(os.path.dirname(__file__), "spec")
L = get_column_letter


def load(name):
    with open(os.path.join(SPEC_DIR, name + ".json")) as fh:
        return json.load(fh)


def cell(spec, coord):
    return spec["cells"].get(coord, {})


def rgb(v):
    return ("FF" + v) if v else None


def font_of(d, name=None, size=None, bold=None, color=None):
    f = d.get("font") or ["Calibri", 11, False, False, None, None]
    return Font(name=name or f[0] or "Calibri", size=size or f[1] or 11,
                bold=f[2] if bold is None else bold, italic=f[3], underline=f[4],
                color=rgb(color if color is not None else f[5]) if (color or f[5]) and (color or f[5]) != "000000" else None)


def side(v):
    return Side(style=v[0], color=rgb(v[1]) if v[1] and v[1] != "000000" else None) if v else Side()


def border_of(d):
    b = d.get("brd") or {}
    return Border(left=side(b.get("left")), right=side(b.get("right")),
                  top=side(b.get("top")), bottom=side(b.get("bottom")))


def align_of(d):
    a = d.get("al")
    if not a:
        return Alignment()
    return Alignment(horizontal=a[0], vertical=a[1], wrap_text=a[2] or None,
                     text_rotation=a[3] or 0, indent=a[4] or 0)


def apply(target, d, font_name=None, fill=True, value=False, nf=True):
    """Style one target cell like reference cell spec d (fonts optionally forced to font_name)."""
    target.font = font_of(d, name=font_name)
    if fill and d.get("fill"):
        target.fill = PatternFill("solid", fgColor=rgb(d["fill"]))
    else:
        target.fill = PatternFill()
    target.border = border_of(d)
    target.alignment = align_of(d)
    if nf and d.get("nf"):
        target.number_format = d["nf"]
    if value and "v" in d:
        target.value = d["v"]


def copy_row(spec, ref_row, ws, row, c1, c2, font_name=None, fill=True, values=False, nf=True):
    for c in range(c1, c2 + 1):
        d = cell(spec, f"{L(c)}{ref_row}")
        apply(ws.cell(row, c), d, font_name, fill, values, nf)
    h = spec["rows"].get(str(ref_row))
    if h and h[0]:
        ws.row_dimensions[row].height = h[0]


def widths(spec, ws, c1, c2, offset=0):
    for c in range(c1, c2 + 1):
        v = spec["cols"].get(L(c))
        if v:
            ws.column_dimensions[L(c + offset)].width = v[0]
            if v[1]:
                ws.column_dimensions[L(c + offset)].hidden = True
            if v[2]:
                ws.column_dimensions[L(c + offset)].outlineLevel = v[2]
        else:
            # a span entry (min..max) may cover this column
            for k, w in spec["cols"].items():
                if w[3] and w[4] and w[3] <= c <= w[4]:
                    ws.column_dimensions[L(c + offset)].width = w[0]
                    break


def page(spec, ws, area=None, titles=None):
    p = spec["print"]
    ps = ws.page_setup
    ps.orientation = p[0]
    ps.paperSize = p[1]
    if p[2] is not None:
        ps.fitToWidth = p[2]
    if p[3] is not None:
        ps.fitToHeight = p[3]
    if p[4]:
        ps.scale = p[4]
    if p[5]:
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    m = spec["margins"]
    pm = ws.page_margins
    pm.left, pm.right, pm.top, pm.bottom, pm.header, pm.footer = m
    if spec.get("footer"):
        l, c, r = spec["footer"]
        if l: ws.oddFooter.left.text = l
        if c: ws.oddFooter.center.text = c
        if r: ws.oddFooter.right.text = r
    if spec.get("zoom"):
        ws.sheet_view.zoomScale = spec["zoom"]
    if spec.get("grid") is False:
        ws.sheet_view.showGridLines = False
    if spec.get("tab"):
        ws.sheet_properties.tabColor = spec["tab"]
    if area:
        ws.print_area = area
    if titles:
        ws.print_title_rows = titles


def fill(v):
    return PatternFill("solid", fgColor=rgb(v)) if v else PatternFill()


def text_rule(ws, rng, word, colour):
    first = rng.split(":")[0].replace("$", "")
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'NOT(ISERROR(SEARCH("{word}",{first})))'], fill=fill(colour)))


def equal_rule(ws, rng, value, colour):
    ws.conditional_formatting.add(rng, CellIsRule(
        operator="equal", formula=[value], fill=fill(colour), font=Font(color=rgb(colour))))
