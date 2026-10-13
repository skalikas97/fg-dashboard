import sys, openpyxl
from openpyxl.utils import get_column_letter as L
path, sheet = sys.argv[1], sys.argv[2]
rows = [int(x) for x in sys.argv[3].split(',')]
maxc = int(sys.argv[4]) if len(sys.argv) > 4 else 10
wb = openpyxl.load_workbook(path)
ws = wb[sheet]
def col(c):
    if c is None: return None
    if c.type == 'rgb': return c.rgb
    if c.type == 'theme': return f'th{c.theme}{c.tint:+.2f}'
    return f'idx{c.indexed}'
def b(s): return '/'.join(f'{k[0]}:{getattr(s,k).style}' for k in ('left','right','top','bottom') if getattr(s,k) and getattr(s,k).style)
for r in rows:
    print(f'-- row {r} h={ws.row_dimensions[r].height} hidden={ws.row_dimensions[r].hidden} lvl={ws.row_dimensions[r].outlineLevel}')
    for c in range(1, maxc+1):
        x = ws.cell(r, c)
        f = x.fill.fgColor if x.fill and x.fill.fill_type == 'solid' else None
        v = repr(x.value)[:40] if x.value is not None else ''
        a = x.alignment
        print(f'  {L(c)}{r} {v} fill={col(f)} font={x.font.name} {x.font.sz} b={x.font.b} col={col(x.font.color)} brd={b(x.border)} al={a.horizontal},{a.vertical},w={a.wrap_text},rot={a.text_rotation} nf={x.number_format!r}')
print('merged', [str(m) for m in ws.merged_cells.ranges][:20])
print('print area', ws.print_area, 'margins', ws.page_margins.left, ws.page_margins.right, ws.page_margins.top, ws.page_margins.bottom, ws.page_margins.header, ws.page_margins.footer)
ps = ws.page_setup; print('setup', ps.orientation, ps.paperSize, ps.fitToWidth, ps.fitToHeight, ps.scale, ws.sheet_properties.pageSetUpPr, ws.print_options.horizontalCentered, ws.print_options.gridLines)
print('view', ws.sheet_view.zoomScale, ws.sheet_view.showGridLines, ws.sheet_view.view, ws.sheet_format.defaultRowHeight, ws.sheet_format.baseColWidth, ws.sheet_properties.outlinePr)
print('dims', {k: (v.width, v.hidden, v.outlineLevel) for k, v in ws.column_dimensions.items()})
