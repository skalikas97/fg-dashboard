"""Render each tab of a workbook to one PNG (scratch copies only; the input file is not changed)."""
import os, subprocess, sys, openpyxl
from openpyxl.worksheet.properties import PageSetupProperties
src, outdir = sys.argv[1], sys.argv[2]
maxrow = int(sys.argv[3]) if len(sys.argv) > 3 else 60
os.makedirs(outdir, exist_ok=True)
names = openpyxl.load_workbook(src).sheetnames
for i, name in enumerate(names):
    wb = openpyxl.load_workbook(src)
    for other in [n for n in wb.sheetnames if n != name]:
        del wb[other]
    ws = wb[name]
    ws.print_area = f"A1:{openpyxl.utils.get_column_letter(min(ws.max_column, 70))}{min(ws.max_row, maxrow)}"
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1; ws.page_setup.orientation = "landscape"
    ws.print_title_rows = None
    base = os.path.join(outdir, f"{i:02d}")
    wb.save(base + ".xlsx")
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", outdir, base + ".xlsx"],
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
    subprocess.run(["pdftoppm", "-f", "1", "-l", "1", "-r", "70", "-png", base + ".pdf", base])
    print(i, name)
